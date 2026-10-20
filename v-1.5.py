import tkinter as tk
from tkinter import ttk
import customtkinter as ctk
from docx import Document
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import tkinter.messagebox as msgbox
from tkinter import filedialog

ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue")

COLOR_OLIVA_40 = "CCCC99"

class AppPlaneaciones(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Planeatrón - Asistente de Planeaciones Didácticas (NEM Semanal)")
        
        # Iniciar maximizada por defecto
        self.after(10, lambda: self.state("zoomed"))
        self.minsize(1000, 650)

        # Configurar estilos nativos ttk para el formulario
        self.style = ttk.Style()
        self.style.theme_use("clam")
        self.style.configure("TLabel", font=("Segoe UI", 10, "bold"), background="#F0F2F5")
        self.style.configure("Header.TLabel", font=("Segoe UI", 12, "bold"), background="#F0F2F5", foreground="#1C3144")
        self.style.configure("TEntry", padding=5)

        # Título Principal
        self.lbl_title = ctk.CTkLabel(
            self, 
            text="Planeatrón - Generador de Planeaciones Semanales", 
            font=ctk.CTkFont(size=20, weight="bold")
        )
        self.lbl_title.pack(pady=(12, 4))

        # Contenedor principal: REBALANCEO DE PROPORCIONES (70% Formulario, 30% Vista previa)
        self.main_container = ctk.CTkFrame(self, fg_color="transparent")
        self.main_container.pack(fill="both", expand=True, padx=15, pady=(0, 10))
        self.main_container.grid_rowconfigure(0, weight=1)
        self.main_container.grid_columnconfigure(0, weight=7, uniform="columnas")  # Más espacio al formulario
        self.main_container.grid_columnconfigure(1, weight=3, uniform="columnas")  # Vista previa compacta

        # Columna izquierda: MARCO DE SCROLL NATIVO (Formulario principal)
        self.scroll_canvas = tk.Canvas(self.main_container, bg="#F0F2F5", highlightthickness=0)
        self.scrollbar = ttk.Scrollbar(self.main_container, orient="vertical", command=self.scroll_canvas.yview)
        
        self.scroll_frame = ttk.Frame(self.scroll_canvas, style="TFrame")
        self.scroll_frame.bind(
            "<Configure>", 
            lambda e: self.scroll_canvas.configure(scrollregion=self.scroll_canvas.bbox("all"))
        )
        
        self.canvas_window = self.scroll_canvas.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        self.scroll_canvas.configure(yscrollcommand=self.scrollbar.set)
        
        self.scroll_canvas.grid(row=0, column=0, sticky="nsew", padx=(0, 0), pady=5)
        self.scrollbar.grid(row=0, column=0, sticky="nse", pady=5)

        self.scroll_canvas.bind('<Configure>', self._on_canvas_configure)
        self.bind_all("<MouseWheel>", self._on_mousewheel)

        # Columna derecha: Vista previa compacta con barra de herramientas de Zoom
        self.preview_col = ctk.CTkFrame(self.main_container, fg_color="transparent")
        self.preview_col.grid(row=0, column=1, sticky="nsew", padx=(12, 0), pady=5)
        self.preview_col.grid_rowconfigure(1, weight=1)
        self.preview_col.grid_columnconfigure(0, weight=1)

        # Encabezado de la Vista Previa + BARRA DE CONTROLES DE ZOOM
        self.header_preview_frame = ctk.CTkFrame(self.preview_col, fg_color="transparent")
        self.header_preview_frame.grid(row=0, column=0, sticky="ew", pady=(0, 5))

        self.lbl_preview = ctk.CTkLabel(
            self.header_preview_frame,
            text="👁 Vista previa",
            font=ctk.CTkFont(size=14, weight="bold"),
            anchor="w"
        )
        self.lbl_preview.pack(side="left")

        # Nivel de zoom inicial (tamaños de fuente)
        self.zoom_level = 10  # Tamaño de letra base en la vista previa
        
        self.btn_zoom_out = ctk.CTkButton(
            self.header_preview_frame, text="🔍-", width=30, height=24,
            command=lambda: self._cambiar_zoom(-1)
        )
        self.btn_zoom_out.pack(side="right", padx=(2, 0))

        self.btn_zoom_reset = ctk.CTkButton(
            self.header_preview_frame, text="100%", width=40, height=24,
            command=self._reset_zoom
        )
        self.btn_zoom_reset.pack(side="right", padx=(2, 0))

        self.btn_zoom_in = ctk.CTkButton(
            self.header_preview_frame, text="🔍+", width=30, height=24,
            command=lambda: self._cambiar_zoom(1)
        )
        self.btn_zoom_in.pack(side="right", padx=(2, 0))

        self.preview_scroll = ctk.CTkScrollableFrame(
            self.preview_col,
            fg_color="#DCE1E5",
            corner_radius=4
        )
        self.preview_scroll.grid(row=1, column=0, sticky="nsew")

        # Hoja simulada de Word
        self.preview_page = ctk.CTkFrame(
            self.preview_scroll,
            fg_color="#FFFFFF",
            corner_radius=2,
            border_width=1,
            border_color="#B9C0C4"
        )
        self.preview_page.pack(fill="x", padx=6, pady=6)

        self.lbl_preview_nota = ctk.CTkLabel(
            self.preview_col,
            text="Usa los botones de zoom si necesitas aumentar la letra.",
            font=ctk.CTkFont(size=10),
            text_color="#5A646B",
            anchor="w"
        )
        self.lbl_preview_nota.grid(row=2, column=0, sticky="ew", pady=(5, 0))

        # ==============================================================
        # ESTRUCTURA DE DATOS EN MEMORIA
        # ==============================================================
        self.dias_semana = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]
        self.dia_activo = "Lunes"
        self.datos_dias = {dia: {
            "campo_formativo": "", "escenario": "", "libro": "", "problematica": "",
            "contenido": "", "pda": "", "tema": "", "aprendizaje": "",
            "inicio": "", "desarrollo": "", "cierre": "", "recursos": "", "evaluacion": ""
        } for dia in self.dias_semana}

        # ==============================================================
        # 1. DATOS GENERALES (SEMANALES)
        # ==============================================================
        self.add_section_header("1. Datos Generales (Semanales)", parent=self.scroll_frame)
        self.txt_titulo = self.add_input_field("Título de la Planeación:", parent=self.scroll_frame)
        self.txt_docente = self.add_input_field("Nombre del Docente:", parent=self.scroll_frame)
        self.txt_escuela = self.add_input_field("Escuela / Institución:", parent=self.scroll_frame)
        self.txt_grado = self.add_input_field("Grado y Grupo:", parent=self.scroll_frame)
        self.txt_periodo = self.add_input_field("Periodo / Semana (ej. del 20 al 24 de Octubre):", parent=self.scroll_frame)

        # ==============================================================
        # 2. SELECTOR DE DÍA
        # ==============================================================
        self.add_section_header("2. Planeación Didáctica por Día", parent=self.scroll_frame)
        
        self.frame_selector = ttk.Frame(self.scroll_frame)
        self.frame_selector.pack(fill="x", pady=(5, 10))

        self.btn_dias = {}
        for dia in self.dias_semana:
            btn = tk.Button(
                self.frame_selector, 
                text=dia, 
                font=("Segoe UI", 10, "bold"),
                bg="#E0E0E0", 
                fg="#000000",
                relief="flat",
                command=lambda d=dia: self._cambiar_dia(d)
            )
            btn.pack(side="left", expand=True, fill="x", padx=2)
            self.btn_dias[dia] = btn

        self._actualizar_estilo_botones_dia()

        # CONTENEDOR DINÁMICO NATIVO
        self.frame_dia_dinamico = ttk.Frame(self.scroll_frame)
        self.frame_dia_dinamico.pack(fill="x", expand=True)

        self._crear_campos_dia_dinamico()

        # BOTÓN GENERAR
        self.btn_generar = tk.Button(
            self.scroll_frame, 
            text="📄 Guardar Planeación Word Semanal (.docx)", 
            command=self.generar_word,
            font=("Segoe UI", 11, "bold"),
            bg="#1E4D2B", 
            fg="#FFFFFF",
            activebackground="#14341D",
            activeforeground="#FFFFFF",
            relief="flat",
            height=2
        )
        self.btn_generar.pack(fill="x", padx=5, pady=(20, 20))

        # ESTRUCTURA DE VISTA PREVIA
        self._pv_labels = {}
        self._construir_estrucutra_preview()

        # EVENTOS DE TECLADO
        self._debounce_job = None
        self._vincular_eventos_globales()

        self.actualizar_preview()

    def _cambiar_zoom(self, delta):
        # Limitar el zoom entre tamaño 8 y 18 para evitar distorsiones extremas
        nuevo_zoom = max(8, min(18, self.zoom_level + delta))
        if nuevo_zoom != self.zoom_level:
            self.zoom_level = nuevo_zoom
            self._aplicar_zoom_a_labels()

    def _reset_zoom(self):
        self.zoom_level = 10
        self._aplicar_zoom_a_labels()

    def _aplicar_zoom_a_labels(self):
        sz_titulo = self.zoom_level + 4
        sz_hdr = self.zoom_level
        sz_body = max(7, self.zoom_level - 2)

        self._pv_labels["titulo"].configure(font=ctk.CTkFont(size=sz_titulo, weight="bold"))

        for key, lbl in self._pv_labels.items():
            if key == "titulo":
                continue
            if "_hdr" in key or key.startswith("gen_"):
                lbl.configure(font=ctk.CTkFont(size=sz_hdr, weight="bold"))
            else:
                lbl.configure(font=ctk.CTkFont(size=sz_body))

        # Actualizar indicador porcentual
        porcentaje = int((self.zoom_level / 10) * 100)
        self.btn_zoom_reset.configure(text=f"{porcentaje}%")

    def _on_canvas_configure(self, event):
        self.scroll_canvas.itemconfig(self.canvas_window, width=event.width)

    def _on_mousewheel(self, event):
        self.scroll_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

    def add_section_header(self, text, parent):
        lbl = ttk.Label(parent, text=text, style="Header.TLabel")
        lbl.pack(fill="x", pady=(10, 4))

    def add_input_field(self, label_text, parent):
        lbl = ttk.Label(parent, text=label_text, style="TLabel")
        lbl.pack(fill="x", pady=(2, 0))
        entry = ttk.Entry(parent, font=("Segoe UI", 10))
        entry.pack(fill="x", pady=(0, 4))
        return entry

    def add_textbox_field(self, label_text, parent):
        lbl = ttk.Label(parent, text=label_text, style="TLabel")
        lbl.pack(fill="x", pady=(2, 0))
        tb = tk.Text(parent, height=3, font=("Segoe UI", 10), undo=True, relief="solid", bd=1)
        tb.pack(fill="x", pady=(0, 4))
        return tb

    def _crear_campos_dia_dinamico(self):
        self.add_section_header("Encuadre Curricular (NEM)", parent=self.frame_dia_dinamico)
        self.txt_campo_formativo = self.add_input_field("Campo Formativo:", parent=self.frame_dia_dinamico)
        self.txt_escenario = self.add_input_field("Escenario:", parent=self.frame_dia_dinamico)
        self.txt_libro = self.add_input_field("Libro de Texto:", parent=self.frame_dia_dinamico)
        self.txt_problematica = self.add_textbox_field("Problemática:", parent=self.frame_dia_dinamico)
        self.txt_contenido = self.add_textbox_field("Contenido:", parent=self.frame_dia_dinamico)
        self.txt_pda = self.add_textbox_field("PDA (Proceso de Desarrollo de Aprendizaje):", parent=self.frame_dia_dinamico)
        self.txt_tema = self.add_input_field("Tema del Día:", parent=self.frame_dia_dinamico)
        self.txt_aprendizaje = self.add_textbox_field("Aprendizaje Esperado / Propósito:", parent=self.frame_dia_dinamico)

        self.add_section_header("Secuencia Didáctica", parent=self.frame_dia_dinamico)
        self.txt_inicio = self.add_textbox_field("Inicio (Enganche / Saberes previos):", parent=self.frame_dia_dinamico)
        self.txt_desarrollo = self.add_textbox_field("Desarrollo (Actividades principales):", parent=self.frame_dia_dinamico)
        self.txt_cierre = self.add_textbox_field("Cierre (Evaluación / Puesta en común):", parent=self.frame_dia_dinamico)

        self.add_section_header("Recursos y Evaluación", parent=self.frame_dia_dinamico)
        self.txt_recursos = self.add_textbox_field("Recursos y Materiales del Día:", parent=self.frame_dia_dinamico)
        self.txt_evaluacion = self.add_textbox_field("Criterios e Instrumentos de Evaluación:", parent=self.frame_dia_dinamico)

        self._map_campos = {
            "campo_formativo": self.txt_campo_formativo,
            "escenario": self.txt_escenario,
            "libro": self.txt_libro,
            "problematica": self.txt_problematica,
            "contenido": self.txt_contenido,
            "pda": self.txt_pda,
            "tema": self.txt_tema,
            "aprendizaje": self.txt_aprendizaje,
            "inicio": self.txt_inicio,
            "desarrollo": self.txt_desarrollo,
            "cierre": self.txt_cierre,
            "recursos": self.txt_recursos,
            "evaluacion": self.txt_evaluacion,
        }

    def _guardar_dia_actual_en_memoria(self):
        d = self.datos_dias[self.dia_activo]
        for key, widget in self._map_campos.items():
            if isinstance(widget, ttk.Entry):
                d[key] = widget.get().strip()
            elif isinstance(widget, tk.Text):
                d[key] = widget.get("1.0", "end-1c").strip()

    def _cargar_dia_desde_memoria(self, nuevo_dia):
        d = self.datos_dias[nuevo_dia]
        for key, widget in self._map_campos.items():
            valor = d[key]
            if isinstance(widget, ttk.Entry):
                widget.delete(0, "end")
                widget.insert(0, valor)
            elif isinstance(widget, tk.Text):
                widget.delete("1.0", "end")
                widget.insert("1.0", valor)

    def _actualizar_estilo_botones_dia(self):
        for dia, btn in self.btn_dias.items():
            if dia == self.dia_activo:
                btn.config(bg="#1C3144", fg="#FFFFFF")
            else:
                btn.config(bg="#E0E0E0", fg="#000000")

    def _cambiar_dia(self, nuevo_dia):
        self._guardar_dia_actual_en_memoria()
        self.dia_activo = nuevo_dia
        self._cargar_dia_desde_memoria(nuevo_dia)
        self._actualizar_estilo_botones_dia()
        self.actualizar_preview()

    def _vincular_eventos_globales(self):
        campos_globales = [self.txt_titulo, self.txt_docente, self.txt_escuela, self.txt_grado, self.txt_periodo]
        for campo in campos_globales:
            campo.bind("<KeyRelease>", self._on_input_change)

        for widget in self._map_campos.values():
            widget.bind("<KeyRelease>", self._on_input_change)

    def _on_input_change(self, event=None):
        if self._debounce_job:
            self.after_cancel(self._debounce_job)
        self._debounce_job = self.after(150, self.actualizar_preview)

    def _leer_campos_totales(self, vacio=""):
        self._guardar_dia_actual_en_memoria()

        def simple(widget):
            return widget.get().strip() or vacio

        return {
            "titulo": simple(self.txt_titulo),
            "docente": simple(self.txt_docente),
            "escuela": simple(self.txt_escuela),
            "grado": simple(self.txt_grado),
            "periodo": simple(self.txt_periodo),
            "dias": self.datos_dias
        }

    def _construir_estrucutra_preview(self):
        margen = ctk.CTkFrame(self.preview_page, fg_color="transparent")
        margen.pack(fill="both", expand=True, padx=10, pady=10)
        margen.grid_columnconfigure(0, weight=1)

        self._pv_labels["titulo"] = ctk.CTkLabel(
            margen, text="", font=ctk.CTkFont(size=14, weight="bold"),
            text_color="#000000", wraplength=280, justify="center"
        )
        self._pv_labels["titulo"].grid(row=0, column=0, sticky="ew", pady=(0, 10))

        m_gen = ctk.CTkFrame(margen, fg_color="#B9C0C4", corner_radius=0)
        m_gen.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        m_gen.grid_columnconfigure(0, weight=1)
        m_gen.grid_columnconfigure(1, weight=1)

        coords = [
            ("gen_docente", 0, 0), ("gen_escuela", 0, 1),
            ("gen_grado", 1, 0), ("gen_periodo", 1, 1),
        ]
        for key, r, c in coords:
            lbl = ctk.CTkLabel(
                m_gen, text="", fg_color="#FFFFFF", text_color="#000000",
                font=ctk.CTkFont(size=9), anchor="w", justify="left", wraplength=130, corner_radius=0
            )
            lbl.grid(row=r, column=c, sticky="nsew", padx=1, pady=1, ipadx=4, ipady=2)
            self._pv_labels[key] = lbl

        row_counter = 2
        for dia in self.dias_semana:
            m_dia = ctk.CTkFrame(margen, fg_color="#B9C0C4", corner_radius=0)
            m_dia.grid(row=row_counter, column=0, sticky="ew", pady=(0, 8))
            m_dia.grid_columnconfigure(0, weight=1)
            m_dia.grid_columnconfigure(1, weight=2)

            lbl_hdr = ctk.CTkLabel(
                m_dia, text=f"📅 DÍA: {dia.upper()}", fg_color="#808000", text_color="#FFFFFF",
                font=ctk.CTkFont(size=9, weight="bold"), anchor="w", corner_radius=0
            )
            lbl_hdr.grid(row=0, column=0, columnspan=2, sticky="nsew", padx=1, pady=1, ipadx=4, ipady=2)
            self._pv_labels[f"{dia}_hdr"] = lbl_hdr

            fields_info = [
                ("Tema / Aprendizaje", f"{dia}_tema_app"),
                ("Campo / Escenario", f"{dia}_campo_esc"),
                ("Libro de Texto", f"{dia}_libro"),
                ("Problemática", f"{dia}_prob"),
                ("Contenido / PDA", f"{dia}_cont_pda"),
                ("Inicio", f"{dia}_inicio"),
                ("Desarrollo", f"{dia}_desarrollo"),
                ("Cierre", f"{dia}_cierre"),
                ("Recursos / Eval.", f"{dia}_rec_eval"),
            ]

            for r_idx, (tit, k_suffix) in enumerate(fields_info, start=1):
                ctk.CTkLabel(
                    m_dia, text=tit, fg_color="#FFFFFF", text_color="#000000",
                    font=ctk.CTkFont(size=8, weight="bold"), anchor="w", corner_radius=0
                ).grid(row=r_idx, column=0, sticky="nsew", padx=1, pady=1, ipadx=3, ipady=2)

                lbl = ctk.CTkLabel(
                    m_dia, text="", fg_color="#FFFFFF", text_color="#000000",
                    font=ctk.CTkFont(size=8), anchor="w", justify="left", wraplength=170, corner_radius=0
                )
                lbl.grid(row=r_idx, column=1, sticky="nsew", padx=1, pady=1, ipadx=3, ipady=2)
                self._pv_labels[k_suffix] = lbl

            row_counter += 1

    def actualizar_preview(self):
        if not self.winfo_exists():
            return

        datos = self._leer_campos_totales(vacio="")

        def v(val):
            return val if val else "(Sin capturar)"

        self._pv_labels["titulo"].configure(text=datos["titulo"] or "PLANEACIÓN DIDÁCTICA SEMANAL")
        self._pv_labels["gen_docente"].configure(text=f"Docente: {v(datos['docente'])}")
        self._pv_labels["gen_escuela"].configure(text=f"Escuela: {v(datos['escuela'])}")
        self._pv_labels["gen_grado"].configure(text=f"Grado/Grupo: {v(datos['grado'])}")
        self._pv_labels["gen_periodo"].configure(text=f"Periodo: {v(datos['periodo'])}")

        for dia in self.dias_semana:
            d = datos["dias"][dia]
            self._pv_labels[f"{dia}_tema_app"].configure(text=f"Tema: {v(d['tema'])}\nPropósito: {v(d['aprendizaje'])}")
            self._pv_labels[f"{dia}_campo_esc"].configure(text=f"{v(d['campo_formativo'])} ({v(d['escenario'])})")
            self._pv_labels[f"{dia}_libro"].configure(text=v(d["libro"]))
            self._pv_labels[f"{dia}_prob"].configure(text=v(d["problematica"]))
            self._pv_labels[f"{dia}_cont_pda"].configure(text=f"Contenido: {v(d['contenido'])}\nPDA: {v(d['pda'])}")
            self._pv_labels[f"{dia}_inicio"].configure(text=v(d["inicio"]))
            self._pv_labels[f"{dia}_desarrollo"].configure(text=v(d["desarrollo"]))
            self._pv_labels[f"{dia}_cierre"].configure(text=v(d["cierre"]))
            self._pv_labels[f"{dia}_rec_eval"].configure(text=f"Recursos: {v(d['recursos'])}\nEval: {v(d['evaluacion'])}")

    # GENERACIÓN DOCX
    def set_cell_background(self, cell, fill_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def escribir_dato(self, cell, etiqueta, valor):
        parrafo = cell.paragraphs[0]
        run_etiqueta = parrafo.add_run(f"{etiqueta}: ")
        run_etiqueta.font.bold = True
        parrafo.add_run(valor or "N/A")

    def aplicar_texto_negro(self, doc):
        def pintar(parrafos):
            for parrafo in parrafos:
                for run in parrafo.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)

        pintar(doc.paragraphs)
        for tabla in doc.tables:
            for fila in tabla.rows:
                for celda in fila.cells:
                    pintar(celda.paragraphs)

    def generar_word(self):
        filepath = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Documento de Word", "*.docx"), ("Todos los archivos", "*.*")],
            title="Guardar Planeación Didáctica Semanal",
            initialfile="Planeacion_Semanal.docx"
        )
        
        if not filepath:
            return

        datos = self._leer_campos_totales(vacio="")

        doc = Document()

        titulo_doc = datos["titulo"] or "PLANEACIÓN DIDÁCTICA SEMANAL"
        title = doc.add_heading(titulo_doc, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        table_gen = doc.add_table(rows=2, cols=2, style="Table Grid")
        table_gen.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        data_gen = [
            ("Docente", datos["docente"], "Escuela", datos["escuela"]),
            ("Grado/Grupo", datos["grado"], "Periodo/Semana", datos["periodo"]),
        ]

        for i, (e_izq, v_izq, e_der, v_der) in enumerate(data_gen):
            celdas = table_gen.rows[i].cells
            self.escribir_dato(celdas[0], e_izq, v_izq)
            self.escribir_dato(celdas[1], e_der, v_der)

        doc.add_paragraph()

        dias_con_contenido = 0

        for dia in self.dias_semana:
            d = datos["dias"][dia]

            tiene_texto = any(bool(v.strip()) for v in d.values())
            if not tiene_texto:
                continue

            dias_con_contenido += 1

            table_dia = doc.add_table(rows=14, cols=2, style="Table Grid")
            table_dia.alignment = WD_TABLE_ALIGNMENT.CENTER

            hdr_dia = table_dia.cell(0, 0).merge(table_dia.cell(0, 1))
            hdr_dia.text = f"DÍA: {dia.upper()}"
            self.set_cell_background(hdr_dia, COLOR_OLIVA_40)
            for p in hdr_dia.paragraphs:
                for r in p.runs:
                    r.font.bold = True

            filas_dia = [
                ("Tema General", d["tema"]),
                ("Campo Formativo", d["campo_formativo"]),
                ("Escenario", d["escenario"]),
                ("Libro de Texto", d["libro"]),
                ("Problemática", d["problematica"]),
                ("Contenido", d["contenido"]),
                ("PDA", d["pda"]),
                ("Aprendizaje Esperado", d["aprendizaje"]),
                ("Inicio", d["inicio"]),
                ("Desarrollo", d["desarrollo"]),
                ("Cierre", d["cierre"]),
                ("Recursos y Materiales", d["recursos"]),
                ("Evaluación", d["evaluacion"]),
            ]

            for idx, (etiqueta, valor) in enumerate(filas_dia, start=1):
                f_cells = table_dia.rows[idx].cells
                
                p_label = f_cells[0].paragraphs[0]
                r_label = p_label.add_run(etiqueta)
                r_label.font.bold = True
                
                f_cells[1].text = valor or "N/A"

            doc.add_paragraph()

        if dias_con_contenido == 0:
            msgbox.showwarning("Atención", "No capturaste información en ningún día de la semana. Se guardó únicamente el encabezado general.")

        self.aplicar_texto_negro(doc)

        try:
            doc.save(filepath)
            msgbox.showinfo("Éxito", f"¡Planeación Semanal guardada correctamente en:\n{filepath}")
        except Exception as e:
            msgbox.showerror("Error", f"No se pudo guardar el archivo.\nDetalle: {e}")

if __name__ == "__main__":
    app = AppPlaneaciones()
    app.mainloop()
